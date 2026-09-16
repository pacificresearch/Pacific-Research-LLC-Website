"""Outputs: final_report.md, property_rankings.csv/.xlsx, market_rankings.xlsx, raw_data.json. Formulas, frozen headers, filters."""
from __future__ import annotations

import csv
import json
import logging
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import Config
from .models import MarketProfile

log = logging.getLogger("prg.report")
CUR, PCT, NUM = '"$"#,##0', "0.00%", "#,##0"


def _fmt_money(v):
    return "UNKNOWN" if v is None else f"${v:,.0f}"


def _fmt_pct(v):
    return "UNKNOWN" if v is None else f"{v:.1%}"


def _sheet(wb: Workbook, title: str, header: List[str], rows: List[List[Any]], formats: Dict[str, str] = None, widths: Dict[str, int] = None):
    ws = wb.create_sheet(title[:31])
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{len(rows) + 1}"
    for i, h in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, max(12, min(40, len(h) + 2)))
        fmt = (formats or {}).get(h)
        if fmt:
            for row in range(2, len(rows) + 2):
                ws.cell(row=row, column=i).number_format = fmt
    return ws


def top_property_rows(results: List[dict], ranking_e: List[dict]) -> List[dict]:
    by_id = {r["listing"]["id"]: r for r in results}
    rows = []
    for i, e in enumerate(ranking_e[:20], 1):
        r = by_id[e["id"]]
        l, b, t, rent = r["listing"], r["base"], r["taxes"], r["rent"]["per_building"]
        sc = r["scenarios"]
        ex = r["expenses"]["lines"]
        agent = l.get("listing_agent") or {}
        rows.append({
            "rank": i, "address": l["formatted_address"], "city_state": f"{l['city']}, {l['state']}", "units": r["units"], "units_status": r["units_status"],
            "asking_price": l["price"], "price_per_unit": b["price_per_unit"], "year_built": l.get("year_built"), "sqft": l.get("square_footage"), "days_on_market": l.get("days_on_market"),
            "current_owner_tax_bill": t["current_tax_bill"]["value"], "post_sale_tax_estimate": t["post_sale_tax_estimate"]["value"], "post_sale_tax_risk": t["post_sale_tax_risk"],
            "base_gross_monthly_rent": rent["conservative"]["base"], "rent_range_low_high": f"{rent['conservative']['low']}–{rent['market_upside']['high']}" if rent["conservative"]["low"] else "UNKNOWN",
            "rent_confidence": rent["conservative"].get("grade"), "vacancy": ex["vacancy"]["amount"], "insurance": ex["insurance"]["amount"], "insurance_status": r["insurance"]["status"],
            "management": ex["management"]["amount"], "maintenance": ex["repairs_maintenance"]["amount"], "capex_reserve": ex["capex_reserve"]["amount"], "utilities": r["expenses"]["utilities_total"],
            "standard_noi": b["standard_noi"], "economic_noi": b["economic_noi"], "standard_cap": b["standard_cap_rate"], "economic_cap": b["economic_cap_rate"],
            "mortgage_amount": b["loan_amount"], "interest_rate_assumed": b["interest_rate_assumed"], "monthly_pi": b["monthly_debt_service"],
            "monthly_conservative_cash_flow": b["monthly_cash_flow"], "cash_on_cash": b["cash_on_cash"], "standard_dscr": b["standard_dscr"], "economic_dscr": b["economic_dscr"],
            "break_even_occupancy": b["break_even_occupancy"], "downside_cash_flow": sc.get("downside", {}).get("monthly_cash_flow"), "severe_cash_flow": sc.get("severe", {}).get("monthly_cash_flow"),
            "hpi_5y_cagr": r["market"]["hpi_5y_cagr"], "market_catalyst_score": r["market"]["catalyst_score"], "hazard_score": r["hazards"]["hazard_score"],
            "underwriting_confidence": r["confidence"]["score"], "minimum_cash_required": b["total_cash_required"], "best_financing_structure": r["structures"].get("recommended_structure"),
            "underwriting_ceiling_price": r["negotiation"].get("underwriting_ceiling_price"), "mls_number": l.get("mls_number"), "listing_agent": f"{agent.get('name', '')} {agent.get('phone', '')} {agent.get('email', '')}".strip(),
            "data_sources": "RentCast listings/records/comps; Census PEP; Zillow ZHVI/ZORI; FHFA HPI; BLS LAUS; config assumptions" + ("; FIXTURE" if r["fixture"] else ""),
            "key_unresolved_diligence": " | ".join(r["diligence_items"][:5]), "ranking_e_score": e["score"], "fixture": r["fixture"],
        })
    return rows


def market_rows(markets: List[MarketProfile], results: List[dict]) -> List[dict]:
    import statistics as st
    rows = []
    for m in markets:
        v = lambda k: (m.metrics[k].v() if k in m.metrics else None)
        rs = [r for r in results if r["market"]["cbsa"] == m.cbsa and r["base"]]
        med = lambda f: (st.median([f(r) for r in rs]) if rs else None)
        cat = m.catalyst_components or {}
        rows.append({"stage1_rank": m.stage1_rank, "market": m.name, "state": m.state, "seed": m.seed, "stage1_score": m.stage1_score,
                     "candidate_listings_underwritten": len(rs), "median_candidate_price": med(lambda r: r["listing"]["price"]),
                     "median_candidate_gross_rent": med(lambda r: r["base"]["gross_monthly_rent"]), "median_tax_burden": med(lambda r: r["expenses"]["lines"]["property_taxes"]["amount"]),
                     "median_economic_cap": med(lambda r: r["base"]["economic_cap_rate"]), "median_conservative_cash_flow": med(lambda r: r["base"]["monthly_cash_flow"]),
                     "median_cash_on_cash": med(lambda r: r["base"]["cash_on_cash"] or 0),
                     "population_2024": v("population"), "population_growth_2020_24": v("population_growth_4y"), "domestic_migration_per_1000": v("domestic_migration_per_1000"),
                     "unemployment_rate": v("unemployment_rate"), "employment_growth_1y": v("employment_growth_1y"), "avg_weekly_wage": v("avg_weekly_wage"), "wage_growth_1y": v("wage_growth_1y"),
                     "typical_home_value_zhvi": v("typical_home_value"), "typical_rent_zori": v("typical_rent"), "rent_to_value_metro": v("rent_to_value_metro"), "rent_growth_1y": v("rent_growth_1y"),
                     "hpi_1y": v("hpi_1y"), "hpi_3y_cagr": v("hpi_3y_cagr"), "hpi_5y_cagr": v("hpi_5y_cagr"), "permits_per_1000_pop": v("permits_per_1000_pop"), "permits_vs_pop_growth": v("permits_vs_pop_growth"),
                     "industry_hhi": v("industry_hhi"), "catalyst_score": m.catalyst_score, "catalyst_coverage": cat.get("coverage"), "energy_boom_watch": cat.get("energy_boom_watch"),
                     "supply_mismatch_signal": cat.get("supply_mismatch_signal"), "major_catalysts": "; ".join(f"{p.get('name')} ({p.get('status')})" for p in (m.research.get("catalysts") or [])) or "NOT RESEARCHED",
                     "major_risks": ", ".join(m.research.get("risks") or []) or "not researched", "relocation_incentive": (m.research.get("incentives") or [{}])[0].get("program", "NOT RESEARCHED") if m.research.get("incentives") else "NOT RESEARCHED",
                     "post_sale_tax_risk_state": None, "market_confidence": round(100 * (0.5 + 0.5 * float(cat.get("confidence") or 0)), 0) if cat else 50})
    return rows


def write_all(cfg: Config, run: Dict[str, Any], markets_selected: List[MarketProfile], universe: List[MarketProfile], stage2: Dict[str, Any], out_dir: Path) -> Dict[str, Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    results = stage2.get("results", [])
    ranks = stage2.get("rankings", {})
    top = top_property_rows(results, ranks.get("E_overall_strategy", []))
    mrows = market_rows(markets_selected, results)
    paths = {}
    # raw json
    paths["raw_data.json"] = out_dir / "raw_data.json"
    with open(paths["raw_data.json"], "w") as f:
        json.dump({"run": run, "markets_selected": [m.model_dump() for m in markets_selected], "universe_top100": [m.model_dump() for m in universe[:100]], "stage2": stage2}, f, indent=1, default=str)
    # csv
    paths["property_rankings.csv"] = out_dir / "property_rankings.csv"
    all_rows = top_property_rows(results, ranks.get("E_overall_strategy", [])[:1000])
    with open(paths["property_rankings.csv"], "w", newline="") as f:
        if all_rows:
            w = csv.DictWriter(f, fieldnames=list(all_rows[0].keys()))
            w.writeheader()
            w.writerows(all_rows)
        else:
            f.write("no properties underwritten (RentCast unavailable or no listings)\n")
    paths["property_rankings.xlsx"] = out_dir / "property_rankings.xlsx"
    _property_workbook(cfg, run, results, ranks, top, mrows, paths["property_rankings.xlsx"])
    paths["market_rankings.xlsx"] = out_dir / "market_rankings.xlsx"
    _market_workbook(mrows, universe, paths["market_rankings.xlsx"])
    paths["final_report.md"] = out_dir / "final_report.md"
    paths["final_report.md"].write_text(final_report_md(cfg, run, markets_selected, universe, stage2, top, mrows))
    return paths


def _property_workbook(cfg, run, results, ranks, top, mrows, path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Executive Summary"
    ws.append(["PRG Property Acquisition Screener — run " + run["run_id"]])
    ws["A1"].font = Font(bold=True, size=14)
    for k, v in run["summary"].items():
        ws.append([k, str(v)])
    ws.append([])
    ws.append(["NOTE", "Interest rate, insurance, expenses are ASSUMPTIONS unless a row's status says VERIFIED/PROVIDER DATA. Fixture rows are synthetic."])
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 110
    if top:
        header = list(top[0].keys())
        fmts = {k: CUR for k in ["asking_price", "price_per_unit", "current_owner_tax_bill", "post_sale_tax_estimate", "base_gross_monthly_rent", "vacancy", "insurance", "management", "maintenance", "capex_reserve", "utilities", "standard_noi", "economic_noi", "mortgage_amount", "monthly_pi", "monthly_conservative_cash_flow", "downside_cash_flow", "severe_cash_flow", "minimum_cash_required", "underwriting_ceiling_price"]}
        fmts.update({k: PCT for k in ["standard_cap", "economic_cap", "interest_rate_assumed", "cash_on_cash", "break_even_occupancy", "hpi_5y_cagr"]})
        rows = [[r[h] for h in header] for r in top]
        ws2 = _sheet(wb, "Top Properties", header, rows, fmts)
        # formulas where practical: recompute caps and price/unit from the sheet's own cells (auditable)
        ci = {h: i + 1 for i, h in enumerate(header)}
        for i in range(2, len(rows) + 2):
            L = get_column_letter
            ws2.cell(row=i, column=ci["price_per_unit"]).value = f"={L(ci['asking_price'])}{i}/{L(ci['units'])}{i}"
            ws2.cell(row=i, column=ci["standard_cap"]).value = f"={L(ci['standard_noi'])}{i}/{L(ci['asking_price'])}{i}"
            ws2.cell(row=i, column=ci["economic_cap"]).value = f"={L(ci['economic_noi'])}{i}/{L(ci['asking_price'])}{i}"
            ws2.cell(row=i, column=ci["monthly_pi"]).value = f"=-PMT({L(ci['interest_rate_assumed'])}{i}/12,{cfg.financing.amortization_years * 12},{L(ci['mortgage_amount'])}{i})"
            ws2.cell(row=i, column=ci["economic_dscr"]).value = f"={L(ci['economic_noi'])}{i}/({L(ci['monthly_pi'])}{i}*12)"
    for title, key, cols in [("Cash Flow Ranking", "A_conservative_cash_flow", None), ("Cash-on-Cash Ranking", "B_cash_on_cash", None),
                             ("Appreciation Ranking", "C_cash_flow_plus_appreciation", None), ("Downside Ranking", "D_downside_protected", None), ("Overall Ranking E", "E_overall_strategy", None)]:
        rr = ranks.get(key, [])
        if rr:
            header = list(rr[0].keys())
            _sheet(wb, title, header, [[r.get(h) for h in header] for r in rr], {"price": CUR, "monthly_cash_flow": CUR, "cash_on_cash": PCT, "severe_cf": CUR, "downside_cf": CUR, "coc_denominator": CUR})
        else:
            _sheet(wb, title, ["note"], [["no underwritten properties this run"]])
    if mrows:
        header = list(mrows[0].keys())
        _sheet(wb, "Market Ranking", header, [[r.get(h) for h in header] for r in mrows], {"median_candidate_price": CUR, "typical_home_value_zhvi": CUR, "typical_rent_zori": CUR, "population_growth_2020_24": PCT, "employment_growth_1y": PCT, "wage_growth_1y": PCT, "rent_to_value_metro": PCT, "rent_growth_1y": PCT, "hpi_1y": PCT, "hpi_3y_cagr": PCT, "hpi_5y_cagr": PCT, "median_economic_cap": PCT, "median_cash_on_cash": PCT, "median_conservative_cash_flow": CUR, "median_candidate_gross_rent": CUR, "median_tax_burden": CUR})
    fin_rows, rc_rows, tax_rows, cat_rows, inc_rows, risk_rows, src_rows = [], [], [], [], [], [], []
    for r in results:
        addr = r["listing"]["formatted_address"]
        for s in r["structures"].get("structures", []):
            fin_rows.append([addr, s["down_payment_pct"], s["down_payment"], s["loan_amount"], s["interest_rate_assumed"], s["monthly_debt_service"], s["monthly_cash_flow"], s["cash_on_cash"], s["economic_dscr"], s["total_cash_required"], s["remaining_capital_vs_max"], s["severe_monthly_cash_flow"], r["structures"].get("recommended_structure") == s["down_payment_pct"]])
        for rs in r["rate_sensitivity"]:
            fin_rows.append([addr, cfg.capital.baseline_down_payment_pct, None, None, rs["rate"], rs["monthly_debt_service"], rs["monthly_cash_flow"], rs["cash_on_cash"], rs["economic_dscr"], None, None, None, "rate sensitivity"])
        for c in r["rent"]["comps_used"]:
            rc_rows.append([addr, c["address"], c["price"], c["bedrooms"], c["bathrooms"], c["square_footage"], c["property_type"], c["distance_miles"], c["days_old"], c["source"]])
        t = r["taxes"]
        tax_rows.append([addr, t["current_tax_bill"]["value"], t["current_tax_bill"]["status"], t["assessed_value"]["value"], t["effective_tax_rate_on_assessed"]["value"], t["post_sale_tax_estimate"]["value"], t["post_sale_tax_estimate"]["status"], t["post_sale_tax_risk"], (t["post_sale_treatment"]["value"] or "")[:200], t["post_sale_treatment"]["url"], t.get("assessor_url") or "UNKNOWN", json.dumps(t.get("tax_history"))])
        for f in r["capex_flags"]:
            risk_rows.append([addr, f["flag"], f["status"], f["note"]])
        risk_rows.append([addr, "physical_hazards", "REFERENCE", ", ".join(r["hazards"]["physical_hazard_flags"]) + f" (hazard score {r['hazards']['hazard_score']}; not an insurance premium)"])
        for line, d in r["expenses"]["lines"].items():
            src_rows.append([addr, line, d["amount"], d["status"], d["note"]])
        src_rows.append([addr, "insurance_estimate", r["insurance"]["value"], r["insurance"]["status"], r["insurance"]["note"]])
        for c in r["confidence"]["deductions"]:
            src_rows.append([addr, "confidence_deduction", -c["points"], "SCORING", c["reason"]])
    seen = set()
    for r in results:
        m = r["market"]
        if m["cbsa"] in seen:
            continue
        seen.add(m["cbsa"])
        for p in (m.get("catalyst") or {}).get("projects", []):
            cat_rows.append([m["name"], p.get("name"), p.get("company"), p.get("investment_usd"), p.get("jobs"), p.get("avg_wage"), p.get("announced"), p.get("start"), p.get("completion"), p.get("status"), p.get("credit"), p.get("source_url")])
        for i in (m.get("incentives") or {}).get("programs", []):
            inc_rows.append([m["name"]] + [i.get(k) for k in ["program", "dollar_value", "cash_or_noncash", "eligibility", "minimum_income", "employment_requirement", "home_purchase_requirement", "residency_requirement", "minimum_stay", "application_deadline", "funding_availability", "source_url", "status"]])
    _sheet(wb, "Financing Scenarios", ["address", "down_pct", "down_payment", "loan", "rate", "monthly_P&I", "monthly_cash_flow", "cash_on_cash", "economic_dscr", "total_cash_required", "remaining_vs_200k", "severe_case_cf", "recommended/note"], fin_rows or [["none"]], {"down_pct": PCT, "down_payment": CUR, "loan": CUR, "rate": PCT, "monthly_P&I": CUR, "monthly_cash_flow": CUR, "cash_on_cash": PCT, "total_cash_required": CUR, "remaining_vs_200k": CUR, "severe_case_cf": CUR})
    _sheet(wb, "Rent Comps", ["subject", "comp_address", "rent", "beds", "baths", "sqft", "type", "distance_mi", "days_old", "source"], rc_rows or [["none"]], {"rent": CUR})
    _sheet(wb, "Taxes", ["address", "current_owner_bill", "bill_status", "assessed_value", "effective_rate_on_assessed", "post_sale_estimate", "estimate_status", "POST_SALE_TAX_RISK", "treatment", "treatment_source", "assessor_url", "tax_history"], tax_rows or [["none"]], {"current_owner_bill": CUR, "assessed_value": CUR, "post_sale_estimate": CUR, "effective_rate_on_assessed": "0.000%"})
    _sheet(wb, "Catalysts", ["market", "project", "company_agency", "investment_usd", "jobs", "avg_wage", "announced", "start", "completion", "status", "credit", "source_url"], cat_rows or [["none on file — research pending (see research/markets/README.md)"]], {"investment_usd": CUR})
    _sheet(wb, "Relocation Incentives", ["market", "program", "dollar_value", "cash_or_noncash", "eligibility", "minimum_income", "employment_requirement", "home_purchase_requirement", "residency_requirement", "minimum_stay", "application_deadline", "funding_availability", "source_url", "status"], inc_rows or [["none on file — research pending"]])
    _sheet(wb, "Risks", ["address", "flag", "status", "note"], risk_rows or [["none"]])
    a_rows = [[k, json.dumps(v, default=str)] for k, v in cfg.raw.items() if k not in ("seed_markets",)]
    _sheet(wb, "Assumptions", ["section", "values (config.yaml)"], a_rows, widths={"values (config.yaml)": 120})
    _sheet(wb, "Sources", ["address", "item", "amount", "status", "note/source"], src_rows or [["none"]], {"amount": CUR}, widths={"note/source": 90})
    wb.save(path)


def _market_workbook(mrows, universe: List[MarketProfile], path: Path):
    wb = Workbook()
    wb.remove(wb.active)
    if mrows:
        header = list(mrows[0].keys())
        _sheet(wb, "Selected Markets", header, [[r.get(h) for h in header] for r in mrows], {"typical_home_value_zhvi": CUR, "typical_rent_zori": CUR, "population_growth_2020_24": PCT, "employment_growth_1y": PCT, "rent_to_value_metro": PCT, "hpi_1y": PCT, "hpi_3y_cagr": PCT, "hpi_5y_cagr": PCT})
    urows = []
    for m in universe:
        v = lambda k: (m.metrics[k].v() if k in m.metrics else None)
        urows.append([m.stage1_rank, m.name, m.state, m.seed, m.stage1_score, v("population"), v("population_growth_4y"), v("domestic_migration_per_1000"), v("typical_home_value"), v("typical_rent"), v("rent_to_value_metro"), v("unemployment_rate"), v("hpi_3y_cagr"), v("permits_per_1000_pop"), json.dumps(m.stage1_components)])
    _sheet(wb, "Universe (Stage 1)", ["rank", "market", "state", "seed", "stage1_score", "population", "pop_growth_2020_24", "dom_migration_per_1000", "zhvi", "zori", "rent_to_value", "unemployment", "hpi_3y_cagr", "permits_per_1000", "components"], urows, {"zhvi": CUR, "zori": CUR, "pop_growth_2020_24": PCT, "rent_to_value": PCT, "hpi_3y_cagr": PCT})
    wb.save(path)


def final_report_md(cfg, run, markets, universe, stage2, top, mrows) -> str:
    L = []
    s = run["summary"]
    L.append(f"# PRG Property Acquisition Screener — {run['run_id']}\n")
    if run.get("fixture_mode"):
        L.append("> ⚠️ **FIXTURE MODE — every property below is SYNTHETIC test data, not a real listing.** RentCast was unavailable; the pipeline ran on labeled fixtures to prove the math and the report plumbing.\n")
    L.append(f"_Generated {run['finished_at']}_. Cadence: weekly. Strategy: 2–4 unit pure cash flow, ≤ ${cfg.capital.price_ceiling:,.0f}, conventional investor financing, VA entitlement preserved (never modelled).\n")
    L.append("## 0. Self-test and data sources\n")
    L.append("| Check | Status |\n|---|---|")
    for k, v in run["selftest"].items():
        L.append(f"| {k} | {v} |")
    L.append("")
    L.append("**Assumptions in force (config.yaml, not quotes):** " + f"{cfg.financing.interest_rate:.2%} / {cfg.financing.amortization_years}-yr, baseline {cfg.capital.baseline_down_payment_pct:.0%} down; vacancy {cfg['income']['vacancy_pct']:.0%}, credit loss {cfg['income']['credit_loss_pct']:.0%}, management {cfg['expenses']['management_pct_of_collected']:.0%} of collected, R&M {cfg['expenses']['repairs_pct_of_gsr']:.0%} GSR, capex {cfg['expenses']['capex_pct_of_gsr']:.0%} GSR (raised for pre-1960/1940 stock). Insurance is a hazard-tiered table ESTIMATE.\n")
    L.append("## 1. Stage 1 — market screen\n")
    L.append(f"{len(universe)} metros screened on public data; {len(markets)} selected for Stage 2 ({sum(m.seed for m in markets)} seeds forced in). Component weights active this run: {run['selftest'].get('stage1_components_active', '')}.\n")
    L.append("| Rank | Market | Score | ZHVI | ZORI | Rent/Value | Pop 20→24 | Dom. mig/1000 | Unemp. | HPI 3y | Permits/1000 | Catalyst | Seed |\n|---|---|---|---|---|---|---|---|---|---|---|---|---|")
    for m in markets:
        v = lambda k: (m.metrics[k].v() if k in m.metrics else None)
        L.append(f"| {m.stage1_rank} | {m.name} | {m.stage1_score} | {_fmt_money(v('typical_home_value'))} | {_fmt_money(v('typical_rent'))} | {_fmt_pct(v('rent_to_value_metro'))} | {_fmt_pct(v('population_growth_4y'))} | {v('domestic_migration_per_1000') if v('domestic_migration_per_1000') is not None else 'UNKNOWN'} | {v('unemployment_rate') if v('unemployment_rate') is not None else 'UNKNOWN'} | {_fmt_pct(v('hpi_3y_cagr'))} | {v('permits_per_1000_pop') if v('permits_per_1000_pop') is not None else 'UNKNOWN'} | {m.catalyst_score if m.catalyst_score is not None else 'UNKNOWN'} | {'✓' if m.seed else ''} |")
    L.append("")
    L.append("Discovered (non-seed) markets in the top of the universe that are NOT on the seed list: " + ", ".join(m.name for m in universe[:15] if not m.seed) + ".\n")
    L.append("## 2. Stage 2 — properties\n")
    L.append(f"Listings pulled: {stage2.get('listings_pulled', 0)} · candidates after early filter: {stage2.get('candidates', 0)} · underwritten: {stage2.get('underwritten', 0)} · skipped for API budget: {stage2.get('skipped_for_budget', 0)}.\n")
    if not stage2.get("results"):
        L.append("**No properties were underwritten.** " + s.get("stage2_note", "") + "\n")
    ranks = stage2.get("rankings", {})
    for title, key, col in [("Ranking A — highest conservative monthly cash flow", "A_conservative_cash_flow", "monthly_cash_flow"), ("Ranking B — highest cash-on-cash", "B_cash_on_cash", "cash_on_cash"),
                            ("Ranking C — cash flow + appreciation asymmetry", "C_cash_flow_plus_appreciation", "asymmetry_score"), ("Ranking D — downside-protected", "D_downside_protected", "downside_score"),
                            ("Ranking E — best overall for the stated strategy (confidence-adjusted; components shown)", "E_overall_strategy", "score")]:
        rr = ranks.get(key, [])
        if not rr:
            continue
        L.append(f"### {title}\n")
        L.append("| # | Address | Units | Price | " + col + " | Confidence |" + (" c_cash_flow | c_capital | c_appreciation | c_downside | c_fundamentals | c_living |" if key.startswith("E") else "") + "\n|---|---|---|---|---|---|" + ("---|---|---|---|---|---|" if key.startswith("E") else ""))
        for i, r in enumerate(rr[:10], 1):
            val = r[col]
            val = _fmt_pct(val) if col == "cash_on_cash" else (_fmt_money(val) if col == "monthly_cash_flow" else val)
            extra = f" {r['c_cash_flow_economics']} | {r['c_capital_efficiency']} | {r['c_appreciation_catalysts']} | {r['c_downside_property_risk']} | {r['c_rent_employment_fundamentals']} | {r['c_relocation_living_economics']} |" if key.startswith("E") else ""
            L.append(f"| {i} | {r['address']} | {r['units']} | {_fmt_money(r['price'])} | {val} | {r['confidence']} |{extra}")
        L.append("")
    if top:
        L.append("## 3. Top properties — full underwriting\n")
        by_id = {r["listing"]["id"]: r for r in stage2["results"]}
        for row in top[:20]:
            r = next(x for x in stage2["results"] if x["listing"]["formatted_address"] == row["address"])
            t, rent, b, sc, neg, st = r["taxes"], r["rent"]["per_building"], r["base"], r["scenarios"], r["negotiation"], r["structures"]
            L.append(f"### #{row['rank']} {row['address']} — {row['units']} units [{r['units_status']}] · {_fmt_money(row['asking_price'])} · built {row['year_built'] or 'UNKNOWN'} · {row['sqft'] or 'UNKNOWN'} sqft · DOM {row['days_on_market'] if row['days_on_market'] is not None else 'UNKNOWN'} · confidence {row['underwriting_confidence']}/100\n")
            L.append(f"- **Rent (per building, monthly):** contractual {rent['contractual']['base'] or 'UNKNOWN'} · conservative {rent['conservative']['low']}–{rent['conservative']['base']}–{rent['conservative']['high']} (grade {rent['conservative'].get('grade')}, {rent['conservative'].get('n_comps', 0)} comps, median {rent['conservative'].get('median_comp_distance')} mi / {rent['conservative'].get('median_comp_age_days')} d) · upside {rent['market_upside']['low']}–{rent['market_upside']['base']}–{rent['market_upside']['high']}. {r['rent']['layout_note']}")
            L.append(f"- **Taxes:** CURRENT OWNER'S bill {_fmt_money(t['current_tax_bill']['value'])} [{t['current_tax_bill']['status']}] on assessed {_fmt_money(t['assessed_value']['value'])} (eff. {_fmt_pct(t['effective_tax_rate_on_assessed']['value'])}) → **POST-PURCHASE estimate {_fmt_money(t['post_sale_tax_estimate']['value'])} [{t['post_sale_tax_estimate']['status']}] · POST-SALE TAX RISK = {t['post_sale_tax_risk']}**. {(t['post_sale_treatment']['value'] or 'treatment UNKNOWN')[:220]} Source: {t['post_sale_treatment']['url'] or 'none'}; assessor: {t.get('assessor_url') or 'UNKNOWN'}")
            L.append(f"- **Insurance:** {_fmt_money(r['insurance']['value'])}/yr [{r['insurance']['status']}] — {r['insurance']['note'][:120]}. Sensitivity CF/mo: " + ", ".join(f"x{x['multiplier']}: {_fmt_money(x['monthly_cash_flow'])}" for x in r["insurance_sensitivity"]) + f". Hazards: {', '.join(r['hazards']['physical_hazard_flags']) or 'none flagged'} (physical, not premium).")
            ex = r["expenses"]
            L.append(f"- **Operating (annual):** GSR {_fmt_money(ex['gsr'])} · EGI {_fmt_money(ex['egi'])} · OpEx {_fmt_money(ex['opex'])} (utilities {_fmt_money(ex['utilities_total'])}, mgmt {_fmt_money(ex['lines']['management']['amount'])}, R&M {_fmt_money(ex['lines']['repairs_maintenance']['amount'])}) · **Standard NOI {_fmt_money(b['standard_noi'])} · Economic NOI {_fmt_money(b['economic_noi'])}** (capex reserve {_fmt_money(ex['capex_reserve'])}) · OER {_fmt_pct(b['operating_expense_ratio'])}")
            L.append(f"- **Returns @ {b['down_payment_pct']:.0%} down / {b['interest_rate_assumed']:.2%} (ASSUMED):** loan {_fmt_money(b['loan_amount'])}, P&I {_fmt_money(b['monthly_debt_service'])}/mo · std cap {_fmt_pct(b['standard_cap_rate'])} · **econ cap {_fmt_pct(b['economic_cap_rate'])}** · **CF {_fmt_money(b['monthly_cash_flow'])}/mo** ({_fmt_money(b['annual_cash_flow'])}/yr) · **CoC {_fmt_pct(b['cash_on_cash'])}** on {_fmt_money(b['coc_denominator'])} (down+closing+rehab+lease-up) · DSCR std {b['standard_dscr']} / econ {b['economic_dscr']} · debt yield {_fmt_pct(b['debt_yield'])} · break-even occ. {_fmt_pct(b['break_even_occupancy'])} · GRM {b['grm']} · rent/price {_fmt_pct(b['rent_to_price'])}")
            L.append(f"- **Cash to close & stabilize:** down {_fmt_money(b['down_payment'])} + closing {_fmt_money(b['closing_costs'])} + rehab {_fmt_money(b['immediate_rehab'])} (range {r['rehab_range']['value']} — inspection budget, not in NOI) + lease-up {_fmt_money(b['lease_up_costs'])} + lender reserves {_fmt_money(b['lender_reserves'])} + operating reserve {_fmt_money(b['operating_reserve'])} = **{_fmt_money(b['total_cash_required'])}** · return on total cash {_fmt_pct(b['return_on_total_cash'])}" + (" · ⚠️ EXCEEDS $200K DOWN" if b["exceeds_max_down"] else ""))
            L.append("- **Stress (CF/mo · econ DSCR):** " + " · ".join(f"{k}: {_fmt_money(v['monthly_cash_flow'])} / {v['economic_dscr']}" for k, v in sc.items()))
            L.append("- **Equity structures:** " + " · ".join(f"{x['down_payment_pct']:.0%}: CF {_fmt_money(x['monthly_cash_flow'])}, CoC {_fmt_pct(x['cash_on_cash'])}, DSCR {x['economic_dscr']}, cash {_fmt_money(x['total_cash_required'])}, left vs $200K {_fmt_money(x['remaining_capital_vs_max'])}" for x in st.get("structures", [])) + f". Max-CF structure {st.get('max_cash_flow_structure')}, max-CoC {st.get('max_coc_structure')}, **recommended {st.get('recommended_structure')}** — {st.get('recommendation_note', '')}")
            L.append("- **Rate sensitivity (CF/mo):** " + ", ".join(f"{x['rate']:.2%}: {_fmt_money(x['monthly_cash_flow'])} (DSCR {x['economic_dscr']})" for x in r["rate_sensitivity"]))
            L.append(f"- **Negotiation:** ask {_fmt_money(neg.get('asking_price'))} · price for 9% econ cap {_fmt_money(neg.get('econ_cap_9'))} · 10% {_fmt_money(neg.get('econ_cap_10'))} · 12% CoC {_fmt_money(neg.get('coc_12'))} · 15% CoC {_fmt_money(neg.get('coc_15'))} · DSCR 1.25 {_fmt_money(neg.get('dscr_1_25'))} · 1.40 {_fmt_money(neg.get('dscr_1_40'))} → **underwriting ceiling {_fmt_money(neg.get('underwriting_ceiling_price'))}** ({_fmt_pct(neg.get('ceiling_vs_ask_pct'))} vs ask). Not an offer recommendation.")
            L.append(f"- **Market:** {r['market']['name']} — Stage 1 #{r['market']['stage1_rank']}, catalyst score {r['market']['catalyst_score'] if r['market']['catalyst_score'] is not None else 'UNKNOWN'} ({(r['market'].get('catalyst') or {}).get('coverage', '')}), HPI 5y CAGR {_fmt_pct(r['market']['hpi_5y_cagr'])}, unemployment {r['market']['unemployment_rate'] if r['market']['unemployment_rate'] is not None else 'UNKNOWN'}. Neighborhood (ZIP): income {r['neighborhood'].get('median_hh_income', {}).get('value') or 'UNKNOWN'}, renter share {r['neighborhood'].get('renter_share', {}).get('value') or 'UNKNOWN'}, crime: {r['neighborhood'].get('crime', {}).get('note', 'UNKNOWN')}")
            L.append(f"- **Regulatory:** {r['regulatory']['coverage']}; rent control: {r['regulatory']['rent_control']}; licensing: {r['regulatory']['rental_licensing']}; eviction timeline: {r['regulatory']['eviction_timeline_days']}")
            L.append(f"- **Relocation / Year-1 personal benefit (never in NOI):** {r['market']['incentives']['status']}; cash total {r['market']['incentives']['year1_personal_economic_benefit_cash'] if r['market']['incentives']['year1_personal_economic_benefit_cash'] is not None else 'UNKNOWN'}. State income tax: {r['market']['living']['state_income_tax']}.")
            L.append("- **Flags:** " + "; ".join(f"{f['flag']} [{f['status']}]" for f in r["capex_flags"][:12]))
            L.append("- **Confidence deductions:** " + "; ".join(f"-{c['points']} {c['reason']}" for c in r["confidence"]["deductions"]))
            L.append("- **Unresolved diligence:** " + " ".join(f"({i}) {x}" for i, x in enumerate(r["diligence_items"], 1)))
            L.append(f"- **Listing:** MLS {row['mls_number'] or 'UNKNOWN'} · agent {row['listing_agent'] or 'UNKNOWN'} · source {row['data_sources']}\n")
    L.append("## 4. Market report (selected markets)\n")
    L.append("| Market | Underwritten | Median price | Median gross rent | Median tax | Median econ cap | Median CF | Median CoC | Pop trend | Emp. trend | Wage | HPI 5y | Permits/1000 | Catalysts | Risks | Incentive | Confidence |\n|---|---|---|---|---|---|---|---|---|---|---|---|---|---|---|---|---|")
    for m in mrows:
        L.append(f"| {m['market']} | {m['candidate_listings_underwritten']} | {_fmt_money(m['median_candidate_price'])} | {_fmt_money(m['median_candidate_gross_rent'])} | {_fmt_money(m['median_tax_burden'])} | {_fmt_pct(m['median_economic_cap'])} | {_fmt_money(m['median_conservative_cash_flow'])} | {_fmt_pct(m['median_cash_on_cash'])} | {_fmt_pct(m['population_growth_2020_24'])} | {_fmt_pct(m['employment_growth_1y'])} | {m['avg_weekly_wage'] or 'UNKNOWN'} | {_fmt_pct(m['hpi_5y_cagr'])} | {m['permits_per_1000_pop'] or 'UNKNOWN'} | {m['major_catalysts'][:60]} | {m['major_risks'][:40]} | {m['relocation_incentive'][:30]} | {m['market_confidence']} |")
    L.append("")
    L.append("## 5. Data-quality problems (biggest first)\n")
    for p in s.get("data_quality_problems", []):
        L.append(f"- {p}")
    L.append("\n## 6. API calls used\n")
    for k, v in s.get("api_calls", {}).items():
        L.append(f"- {k}: {v}")
    L.append("\n## 7. Keys / data that would most improve accuracy\n")
    for p in s.get("improvements", []):
        L.append(f"- {p}")
    L.append("\n---\n_No number in this report is a lender quote, an insurance quote, or a verified tax bill unless its status says VERIFIED. UNKNOWN is deliberate._\n")
    return "\n".join(L)
